from __future__ import annotations

from pathlib import Path
from copy import copy
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent
SOURCE = ROOT / "新数值.xlsx"

BLUE = "1F4E78"
TEAL = "0F766E"
WHITE = "FFFFFF"
LIGHT = "EAF2F8"
PALE_TEAL = "DDF3EF"
RED = "FCE8E6"
AMBER = "FFF2CC"
GREEN = "E2F0D9"
GRAY = "E7E6E6"
PURPLE = "E4DFEC"
THIN = Side(style="thin", color="B7C9D6")


RULE_GAPS = [
    ("R-001", "触发循环", "致命", "未定义间接触发能否再次触发被动，也没有单事件重入、每回合次数或链深上限。", "黑玫瑰治疗→沼泽鼠施毒→黑玫瑰再次治疗。", "建立事件ID、来源、直接/间接标签；默认同一效果对同一事件只响应一次。", "阻塞制作"),
    ("R-002", "行动时序", "致命", "未定义回合开始、移动、攻击、技能、回合结束、死亡检查之间的统一顺序。", "回合开始、倒计时、剧毒、再生同时存在。", "给出唯一时序表，并为每个触发点分配阶段编号。", "阻塞制作"),
    ("R-003", "护盾/护甲", "高", "护盾与护甲混用，无法判断是同一资源、不同资源还是旧术语。", "披甲写护盾；蜕壳投掷和岩豚写护甲。", "二选一统一，或明确两者减伤顺序、上限与清除规则。", "需作者确认"),
    ("R-004", "弹药例外", "高", "通用规则说弹药归零后不再触发，但黄金飞梭允许无弹药触发。", "黄金飞梭：无论是否有弹药。", "把例外写成规则字段：强制使用是否消耗弹药、0弹药能否响应。", "需作者确认"),
    ("R-005", "随机与并列", "高", "随机目标和最高/最低目标均无并列及确定性规则。", "随机灵兽、攻击最高、生命最低。", "规定候选集、并列顺序、随机种子和同次多选是否放回。", "阻塞测试"),
    ("R-006", "伤害公式", "高", "缺少攻击、技能攻击、防御、战斗伤害和固定伤害的统一结算公式。", "攻击6+、10伤害、技能伤害+2同时出现。", "写成可计算公式并提供3个手算例。", "阻塞平衡"),
    ("R-007", "暴击边界", "高", "暴击被定义为放大几乎所有数值，但未说明固定伤害、反伤、持续伤害和间接效果。", "暴击词条覆盖伤害、治疗、护盾和状态。", "列出可暴击效果白名单；间接效果默认不可暴击。", "需作者确认"),
    ("R-008", "状态结算", "高", "灼烧、剧毒、再生、亢奋、衰弱的层数消耗点和同阶段先后不完整。", "剧毒与再生都可能在回合结束结算。", "为每个状态记录触发阶段、衰减阶段、叠加上限和来源。", "阻塞测试"),
    ("R-009", "倒计时", "高", "倒计时建立、减到0、触发、重置的时点不统一。", "技能使用减一与敌方释放技能减一并存。", "统一为计数器状态机，明确同事件能否连续减多次。", "需作者确认"),
    ("R-010", "充能/爆能", "高", "能量不足时获得能量的规则，与消耗能量后再倒计时的技能缺少状态转换。", "元气弹、吐纳术、爆能词条。", "绘制0→充能→可释放→消耗→结算的状态机。", "需作者确认"),
    ("R-011", "目标与距离", "高", "第一格、一步范围、三步范围、直线最远等语句缺少统一几何定义。", "正前方一格/第一格、左右相邻格。", "提供棋盘坐标、朝向、距离算法、空格穿透和多目标排序。", "阻塞制作"),
    ("R-012", "移动与控制", "中", "禁足、牵引、戏法等位移效果未说明目标格被占、越界和死亡单位处理。", "禁足、牵引、交换位置。", "制定移动失败、换位、推动、拉取的优先级。", "需补规则"),
    ("R-013", "死亡窗口", "高", "死亡、亡语、复苏、持续光环失效的结算窗口未定义。", "在场存活时、死亡后触发、原位置复活。", "每次伤害后统一进行死亡检查，再按队列处理亡语与复苏。", "阻塞制作"),
    ("R-014", "成长上限", "高", "大量永久或整场成长没有上限，和多重触发、暴击组合可能指数膨胀。", "百战、龙牙斩、角蛙、魔性飞剑。", "设单次、每回合和整场上限，并建立标准战斗时长。", "阻塞平衡"),
    ("R-015", "品质预算", "高", "数值多为翻倍增长，但没有品质功率预算、稀有度和获取成本。", "生命、攻击及多数效果按品质翻倍。", "定义青铜=1.0功率单位及白银/黄金/钻石预算。", "阻塞平衡"),
    ("R-016", "技能长度", "中", "短篇、中篇、长篇可能对应体积，但表中没有统一占格字段和收益预算。", "词条列混合长度与类别。", "拆为size字段，并规定1/2/3格的功率倍率。", "需补字段"),
    ("R-017", "升级继承", "中", "升级行大量留空，无法确认是继承上一品质、保持不变还是待填写。", "黄金/钻石行常只填一个变化字段。", "配置层必须展开为完整值；展示层可折叠差异。", "阻塞导出"),
    ("R-018", "跨战斗成长", "中", "永久成长何时保存、复制、出售、升级后如何处理未定义。", "百战、价值成长。", "定义局内永久状态的生命周期和升级迁移规则。", "需作者确认"),
    ("R-019", "经济系统", "高", "价值、购买、出售、经验存在，但没有价格、商店概率和收益基准。", "吞金鼠、宗门大阵、见闻兽。", "补充经济基准表及每回合期望资源。", "阻塞平衡"),
    ("R-020", "日志与验收", "中", "没有要求战斗日志记录触发源、目标、数值前后和事件链。", "复杂触发图无法仅靠结果排错。", "日志至少记录时间/阶段、来源、效果、目标、前值、变化、后值、事件ID。", "阻塞QA"),
]


BUILD_GUIDE = [
    ("多武器", "通过高频武技、暴击和技能伤害成长滚雪球", "蚁王、角蛙、鬼蜂、饿狼、巨鳄蚁、圣象甲虫", "啃咬、利爪、双镰、乱抓、龙牙斩", "武技触发→额外攻击/成长→暴击→全队武技增伤", "启动快、组合数量多", "成长乘法过强；依赖携带多武技", "封刃、减攻、击杀核心光环", "龙牙斩和高频暴击的成长上限未定义"),
    ("灼烧", "叠加灼烧并利用点燃引爆、治疗或额外伤害", "花椒蟹、白蔷薇、赤精鱼", "引火、鞭炮、火球、妒火、火中取栗", "点燃→灼烧→受击触发伤害→治疗/成长", "持续伤害与续航兼备", "依赖状态留存；结算顺序敏感", "净化、快速击杀、状态免疫", "点燃、灼烧和攻击触发的关系需统一"),
    ("剧毒", "通过淬毒施加剧毒，再在回合末或引爆时结算", "蕈章、沼泽鼠、黑玫瑰", "毒钩、毒雾、蜈蚣锁、泰诺地龙、毒腺", "治疗/施毒→剧毒→治疗或护盾→再次触发", "持续输出、可和治疗生存交叉", "存在明显循环风险；启动依赖目标", "净化、爆发、禁止治疗", "黑玫瑰与沼泽鼠需要重入保护"),
    ("护盾", "不断获得护盾并转化为团队生存或技能成长", "岩豚、泥沼妖、圣象甲虫、鲛人枪手、巨伞蕈", "披甲、叠甲、凝血铠甲、双盾、膨胀力场", "武技/增减益→护盾→分发护盾/技能增伤", "容错高、保护核心", "护盾与护甲定义混乱；可能无限叠加", "贯通、破盾、持续伤害", "先统一护盾与护甲，再谈平衡"),
    ("弹药", "用有限爆发换取前期优势，再通过装填维持输出", "铁丸子、玄铁龟", "袖箭、爆破符、飞针盒、双镰、快速换弹、黄金飞梭", "消耗弹药→装填→多重触发/暴击→额外使用", "爆发强，技能顺序有决策", "0弹药例外冲突；装填收益缺少预算", "拖长战斗、打断装填", "黄金飞梭的无弹药触发必须单独定规"),
    ("充能", "积累能量，在爆能门槛释放高价值效果", "烁德童子、赤精鱼、玄铁龟", "蓄能、吐纳术、轻击、剑气、天火咒、元气弹", "技能充能→达到门槛→爆能→范围输出/团队增益", "节奏清晰，适合形成独立身份", "能量不足和倒计时状态转换不清", "窃取能量、延迟、快速压血", "最有潜力成为区别于大巴扎的原创支柱"),
    ("暴击", "提高暴击率，用暴击驱动额外伤害和全队成长", "巨鳄蚁、见闻兽", "凌厉、专注、激昂、龙牙斩、黄金飞梭", "友方行动→暴击率→暴击→增伤/额外使用", "爆发反馈强，和多武器自然结合", "暴击适用范围过宽；容易乘法失控", "降暴击、减少行动频率、击杀暴击核心", "先建立暴击白名单和成长上限"),
    ("增减益", "利用亢奋与衰弱放大己方成长并压制敌方", "黑犬、蚁王", "巨力、尾鞭、风刃、全力冲撞、气场", "施加亢奋/衰弱→成长技能/护盾→继续施加", "能连接多个体系", "翻倍/减半规则对结算顺序极敏感", "净化、免疫、优先击杀辅助", "需要定义先加层还是先放大数值"),
]


BAZAAR_AUDIT = [
    ("袖箭", "左轮手枪", "技能原型、弹药爆发", "高", "改为灵兽携带并加入格子射程", "中高", "保留弹药定位，重做核心触发或装填决策"),
    ("爆破符", "手雷", "单发弹药、低生命目标", "高", "改为符箓和三步范围", "中高", "加入延迟爆破、范围占位或可拆除机制"),
    ("钝化", "流星索", "弹药攻击并降低攻击", "高", "加入格子范围", "高", "把减攻改为牵引、缠绕或阵型破坏"),
    ("冲刺拳", "靴里剑", "首次使用必暴击", "高", "使用武技与直线目标", "高", "让冲刺产生真实位移和位置风险"),
    ("飞针盒", "手里剑", "弹药转多重触发", "高", "改名与目标规则", "高", "改成散射、标记或回收飞针循环"),
    ("装填弹药", "火药角", "给相邻/下一个物品装填", "高", "技能化", "高", "改为灵兽能量分配或回合准备动作"),
    ("双镰", "双管霰弹枪", "弹药、多重触发", "高", "近战题材与直线目标", "高", "加入命中两个不同格或回旋返程"),
    ("快速换弹", "填弹杆", "为相邻物品补满弹药", "高", "增加暴击率", "高", "改为选择性装填并产生过热/空窗成本"),
    ("凝血铠甲", "木桶", "其他武技触发护盾成长", "中高", "绑定灵兽与术法", "中高", "强化受伤、流血或牺牲生命的独特转换"),
    ("魔性飞剑", "步枪", "有限弹药、自身伤害成长", "中高", "改为飞剑主题", "中高", "加入飞剑离场、召回和占据棋盘空间"),
    ("横扫", "武士刀", "范围武器", "中", "明确格子范围", "中", "让朝向、站位和命中数量决定收益"),
    ("百战", "狼筅", "战后永久成长", "中高", "胜利额外成长", "中高", "增加败北保留、击杀记录或流派抉择"),
    ("双盾", "珊瑚甲", "购买特定类型后永久提高护盾", "高", "目标变为两名灵兽", "高", "换成编队防线或相邻单位共同承伤"),
    ("尾鞭", "十手", "施加减益时武器成长", "高", "衰弱和格子目标", "中高", "让尾鞭改变朝向/拉动敌人形成原创空间玩法"),
    ("啃食", "吸血章鱼", "吸血且随暴击成长", "高", "灵兽武技化", "高", "改为吞噬状态、尸体或临时属性"),
    ("风刃", "盐钳海盗", "施加增减益后伤害成长", "高", "格子射程", "高", "让风刃沿路径传播并受站位影响"),
    ("乱抓", "蝴蝶刀", "低基础、多段武器", "中高", "灵兽主题", "中高", "加入随机格连击或未命中补偿"),
    ("钩链", "锁镰", "相邻技能成长与暴击", "高", "以衰弱触发", "高", "转为真实拉扯敌我位置"),
    ("锋锐鳞片", "带刃悬浮板", "相邻武技触发额外伤害", "高", "换题材", "高", "改为受击反射或鳞片消耗资源"),
    ("全力冲撞", "标枪", "有限弹药、直线爆发、团队增益", "中高", "亢奋体系", "中高", "加入冲撞位移和撞墙收益"),
    ("龙牙斩", "赛博铁尺", "暴击使全队武器成长", "高", "武技主题", "高", "限定同列/同灵兽或消耗暴击标记"),
    ("膨胀力场", "查理森先生", "增益触发护盾成长", "中高", "范围友方护盾", "中", "强化棋盘范围变化与站位选择"),
    ("黄金飞梭", "投掷飞刀", "他物暴击时额外使用", "高", "加入弹药但又绕过弹药", "高", "重做为可回收的场上投射物"),
    ("暴起", "电鳗", "敌方使用进攻物品会加速大招", "高", "衰弱与格子范围", "高", "改为蓄势姿态，可被移动或控制打断"),
    ("倾泻", "火炮阵列", "其他武技加速范围多段攻击", "高", "灵兽技能化", "高", "加入炮口方向、友伤或装填阵位"),
    ("圣象甲虫", "船首像", "在场光环强化其他单位和自身武器", "中高", "转为灵兽", "中高", "加入前排承伤或队形完整度条件"),
    ("吞金鼠", "龙涎香", "购买特定类别提高价值并提供治疗", "中高", "改成灵兽回合治疗", "中", "绑定吞噬库存、占位或风险回报"),
    ("见闻兽", "通缉海报", "战斗胜利提供额外经验", "中高", "加入站场暴击光环", "中高", "让侦察信息改变路线或敌方预览"),
]


def setup_sheet(ws, title: str, headers: list[str], widths: dict[int, float] | None = None):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    for c, value in enumerate(headers, 1):
        cell = ws.cell(1, c, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="STHeiti", color=WHITE, bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=TEAL))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 16)


def style_body(ws, max_row: int, max_col: int):
    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.font = Font(name="STHeiti", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"


def risk_for_skill(name: str, effect: str, ability: str, target: str) -> tuple[str, str]:
    text = " ".join([name or "", effect or "", ability or "", target or ""])
    items = []
    severity = "低"
    if "随机" in text or "最高" in text or "最低" in text:
        items.append("随机/并列规则")
        severity = "中"
    if "倒计时" in text:
        items.append("倒计时时点")
        severity = "中"
    if "暴击" in text:
        items.append("暴击边界")
        severity = "中"
    if "永久" in text or "每当" in text or "获得伤害" in text or "伤害提升" in text:
        items.append("成长上限")
        severity = "中"
    if name == "黄金飞梭":
        items.append("与弹药通则冲突")
        severity = "高"
    if name in {"黑玫瑰", "沼泽鼠"}:
        items.append("循环触发")
        severity = "致命"
    if "护甲" in text:
        items.append("护甲/护盾术语")
        severity = "高"
    return severity, "；".join(dict.fromkeys(items)) or "待标准对局验证"


def normalized_skill_rows(src):
    specs = [("青铜技能", 4, "B"), ("白银技能", 3, "S"), ("黄金+技能", 2, "G")]
    out = []
    global_index = 0
    for sheet_name, group_size, prefix in specs:
        ws = src[sheet_name]
        base_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
        for r in base_rows:
            global_index += 1
            base = [ws.cell(r, c).value for c in range(1, 13)]
            for offset in range(group_size):
                rr = r + offset
                if rr > ws.max_row:
                    continue
                vals = [ws.cell(rr, c).value for c in range(1, 13)]
                rarity = vals[1]
                if not rarity:
                    continue
                effective = [vals[c] if vals[c] not in (None, "") else base[c] for c in range(12)]
                name = base[0]
                ability = str(effective[2] or "")
                target = str(effective[4] or "")
                effect = str(effective[5] or "")
                sev, risk = risk_for_skill(name, effect, ability, target)
                out.append([
                    f"SK-{prefix}{global_index:03d}-{offset+1}", name, sheet_name, rr, rarity,
                    ability, effective[3], target, effect, effective[6], effective[7], effective[8],
                    effective[9], effective[10], sev, risk, "展示层已继承空白字段；制作配置仍需作者确认"
                ])
    return out


def normalized_pet_rows(src):
    specs = [("宠物", 4, "B"), ("白银宠物", 3, "S")]
    out = []
    global_index = 0
    for sheet_name, group_size, prefix in specs:
        ws = src[sheet_name]
        base_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value]
        for r in base_rows:
            global_index += 1
            base = [ws.cell(r, c).value for c in range(1, 13)]
            for offset in range(group_size):
                rr = r + offset
                if rr > ws.max_row:
                    continue
                vals = [ws.cell(rr, c).value for c in range(1, 13)]
                rarity = vals[2]
                if not rarity:
                    continue
                effective = [vals[c] if vals[c] not in (None, "") else base[c] for c in range(12)]
                name = base[1]
                effect = str(effective[7] or "")
                sev, risk = risk_for_skill(name, effect, "", "")
                out.append([
                    f"PET-{prefix}{global_index:03d}-{offset+1}", name, sheet_name, rr, rarity,
                    effective[0], effective[3], effective[4], effective[5], effective[6], effect,
                    effective[8], effective[9], effective[10], effective[11], sev, risk,
                    "展示层已继承空白字段；制作配置仍需作者确认"
                ])
    return out


def build_main_workbook():
    src = load_workbook(SOURCE, data_only=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "00_先看这里"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 92
    ws.merge_cells("A1:B1")
    ws["A1"] = "灵兽构筑战斗系统｜评审阅读版"
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="STHeiti", color=WHITE, bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 42
    overview = [
        ("用途", "帮助非策划读者理解系统、构筑关系和制作风险；不替作者修改原始数值。"),
        ("建议阅读顺序", "00先看这里 → 04构筑导航 → 01技能总表/02灵兽总表 → 05规则风险 → 06大巴扎审计。"),
        ("系统一句话", "灵兽携带不同长度和类别的技能，在有朝向与距离的棋盘上，通过状态、资源和触发链形成构筑。"),
        ("核心对象", "灵兽：生命/攻击/防御与被动载体；技能：主动或触发效果；机制：状态、资源及结算规则。"),
        ("品质", "青铜→白银→黄金→钻石。原表升级行的空白在本阅读版中按基础行继承，仅用于阅读，不代表正式配置已确认。"),
        ("两层元素", "点燃负责施加/引爆灼烧；淬毒负责施加/触发剧毒。这个双层结构是值得强化的原创方向。"),
        ("主要构筑", "多武器、灼烧、剧毒、护盾、弹药、充能、暴击、增减益。"),
        ("当前最大风险", "没有统一行动时序与触发重入规则；复杂联动可能产生循环，数值也缺乏功率预算和对局验证。"),
        ("如何读风险", "致命=无法可靠制作；高=会影响规则或平衡；中=需要补字段/边界；低=主要靠标准对局验证。"),
        ("原创性判断", "构筑表达和多个具体原型与《The Bazaar》高度接近；格子射程、灵兽载体、移动与双层元素可发展为核心差异。"),
        ("数据规模", "71个技能原型、25个灵兽原型、约50个已命名机制；数量足够做原型，但不等于已经平衡。"),
        ("原件保护", "来源：新数值.xlsx、交互关系.pdf。原文件未修改。"),
    ]
    for i, (k, v) in enumerate(overview, 3):
        ws.cell(i, 1, k).font = Font(name="STHeiti", bold=True, color=BLUE)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(i, 2, v).font = Font(name="STHeiti", size=10)
        for c in (1, 2):
            ws.cell(i, c).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(i, c).border = Border(bottom=THIN)
        ws.row_dimensions[i].height = 35

    skill_headers = ["建议ID", "技能名", "来源表", "原行", "品质", "能力", "防御", "射程/目标", "效果", "词条", "主要配合对象", "定位", "套路", "参考原型", "风险等级", "阅读提示", "继承说明"]
    ws = wb.create_sheet("01_技能总表")
    setup_sheet(ws, ws.title, skill_headers, {1: 18, 2: 16, 3: 14, 4: 8, 5: 10, 6: 18, 7: 10, 8: 32, 9: 58, 10: 24, 11: 22, 12: 28, 13: 20, 14: 20, 15: 12, 16: 32, 17: 28})
    rows = normalized_skill_rows(src)
    for row in rows:
        ws.append(row)
    style_body(ws, ws.max_row, len(skill_headers))

    pet_headers = ["建议ID", "灵兽名", "来源表", "原行", "品质", "套路", "词条", "HP", "攻击", "防御", "效果", "主要配合对象", "升级理由", "套路定位", "参考原型", "风险等级", "阅读提示", "继承说明"]
    ws = wb.create_sheet("02_灵兽总表")
    setup_sheet(ws, ws.title, pet_headers, {1: 18, 2: 16, 3: 14, 4: 8, 5: 10, 6: 18, 7: 15, 8: 10, 9: 10, 10: 10, 11: 58, 12: 22, 13: 28, 14: 28, 15: 20, 16: 12, 17: 32, 18: 28})
    rows = normalized_pet_rows(src)
    for row in rows:
        ws.append(row)
    style_body(ws, ws.max_row, len(pet_headers))

    ws = wb.create_sheet("03_机制词典")
    src_ws = src["词条"]
    headers = [src_ws.cell(1, c).value or f"原列{c}" for c in range(1, 10)]
    headers += ["评审备注"]
    setup_sheet(ws, ws.title, headers, {1: 12, 2: 16, 3: 18, 4: 18, 5: 24, 6: 8, 7: 24, 8: 8, 9: 70, 10: 32})
    for r in range(2, src_ws.max_row + 1):
        vals = [src_ws.cell(r, c).value for c in range(1, 10)]
        note = "机制ID为空，建议补稳定英文/机器ID" if vals[2] else "未定义占位，不能进入制作配置"
        ws.append(vals + [note])
    style_body(ws, ws.max_row, len(headers))

    ws = wb.create_sheet("04_构筑导航")
    headers = ["构筑", "目标", "核心灵兽", "核心技能", "运转链", "优点", "缺点", "反制", "评审重点"]
    setup_sheet(ws, ws.title, headers, {1: 15, 2: 32, 3: 32, 4: 38, 5: 48, 6: 26, 7: 30, 8: 26, 9: 38})
    for row in BUILD_GUIDE:
        ws.append(row)
    style_body(ws, ws.max_row, len(headers))

    ws = wb.create_sheet("05_规则风险")
    headers = ["编号", "主题", "严重度", "缺口", "证据/示例", "建议规则", "影响"]
    setup_sheet(ws, ws.title, headers, {1: 12, 2: 18, 3: 12, 4: 56, 5: 48, 6: 58, 7: 16})
    for row in RULE_GAPS:
        ws.append(row)
    style_body(ws, ws.max_row, len(headers))

    ws = wb.create_sheet("06_大巴扎审计")
    headers = ["当前内容", "表内参考原型", "相似核心", "借鉴程度", "已做改造", "辨识度风险", "建议改造方向"]
    setup_sheet(ws, ws.title, headers, {1: 18, 2: 22, 3: 40, 4: 14, 5: 38, 6: 16, 7: 54})
    for row in BAZAAR_AUDIT:
        ws.append(row)
    style_body(ws, ws.max_row, len(headers))

    ws = wb.create_sheet("07_原始数值模版")
    src_ws = src["模版"]
    for row in src_ws.iter_rows():
        for cell in row:
            ws[cell.coordinate] = cell.value
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        if cell.value is not None:
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="STHeiti", color=WHITE, bold=True)
            cell.alignment = Alignment(wrap_text=True, horizontal="center")
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14

    for name in ["01_技能总表", "02_灵兽总表", "05_规则风险", "06_大巴扎审计"]:
        ws = wb[name]
        severity_col = None
        for c in range(1, ws.max_column + 1):
            if ws.cell(1, c).value in {"风险等级", "严重度", "辨识度风险"}:
                severity_col = c
                break
        if severity_col:
            letter = get_column_letter(severity_col)
            for word, color in [("致命", RED), ("高", RED), ("中高", AMBER), ("中", AMBER), ("低", GREEN)]:
                ws.conditional_formatting.add(
                    f"{letter}2:{letter}{ws.max_row}",
                    FormulaRule(formula=[f'{letter}2="{word}"'], fill=PatternFill("solid", fgColor=color)),
                )

    path = OUT / "新数值_评审阅读版.xlsx"
    wb.save(path)
    return path


def build_simple_workbook(filename: str, sheet_name: str, headers: list[str], rows: list[tuple], widths: dict[int, float]):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    setup_sheet(ws, sheet_name, headers, widths)
    for row in rows:
        ws.append(row)
    style_body(ws, ws.max_row, len(headers))
    path = OUT / filename
    wb.save(path)
    return path


if __name__ == "__main__":
    files = [build_main_workbook()]
    files.append(build_simple_workbook(
        "大巴扎借鉴审计.xlsx", "借鉴审计",
        ["当前内容", "表内参考原型", "相似核心", "借鉴程度", "已做改造", "辨识度风险", "建议改造方向"],
        BAZAAR_AUDIT, {1: 18, 2: 22, 3: 40, 4: 14, 5: 38, 6: 16, 7: 54}
    ))
    files.append(build_simple_workbook(
        "规则缺口清单.xlsx", "规则缺口",
        ["编号", "主题", "严重度", "缺口", "证据/示例", "建议规则", "影响"],
        RULE_GAPS, {1: 12, 2: 18, 3: 12, 4: 56, 5: 48, 6: 58, 7: 16}
    ))
    print("\n".join(str(p) for p in files))
